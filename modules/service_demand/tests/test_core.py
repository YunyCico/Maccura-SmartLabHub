"""回归测试：ZIP 解析 + 参数校验 + Token 缓存。

全部离线可跑（不联网）：
    .venv/Scripts/python.exe -m pytest tests -q
"""
from __future__ import annotations

import io
import json
import sys
import time
import zipfile
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "backend"))

from app.data_parser import (  # noqa: E402
    PlatformParseError, flatten_metric_payload, parse_zip, profile_column,
)
from app.main import ConditionItem, PassRateRequest, PullRequest  # noqa: E402
from app.platform_client import (  # noqa: E402
    ZIP_TTL_SECONDS, DataPlatformClient, TokenCache, _classify,
)


def make_zip(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries.items():
            z.writestr(name, data)
    return buf.getvalue()


# ---------- ZIP 目录结构（文档 8.2） ----------

def test_parse_document_zip_structure():
    """严格按文档结构：tables/ metrics/ combinations/ manifest.json"""
    csv = "MATNR,LBKUM\n100001,1\n100002,2\n"
    cols = json.dumps({
        "tableName": "table1",
        "chineseName": "物料库存表",
        "columns": [
            {"fieldName": "MATNR", "comment": "物料编号",
             "fieldType": "VARCHAR(50)", "length": "50"},
            {"fieldName": "LBKUM", "comment": "批次数量",
             "fieldType": "INT", "length": ""},
        ],
    }, ensure_ascii=False)
    z = make_zip({
        "tables/table1.columns.json": cols.encode("utf-8"),
        "tables/table1.csv": csv.encode("utf-8-sig"),
        "metrics/sales_amount.json": b'[{"a":1,"b":"x"}]',
        "combinations/sales_dashboard.json": b'{"rows":[{"c":2}]}',
        "manifest.json": json.dumps({
            "tables": [{"tableName": "table1", "chineseName": "物料库存表",
                        "columnsFile": "tables/table1.columns.json",
                        "dataFile": "tables/table1.csv", "rowCount": 2}],
            "metrics": [{"metricCode": "sales_amount"}],
            "combinations": [{"combinationCode": "sales_dashboard"}],
            "message": "表: 请求 1, 成功返回 1 | 指标: 请求 1, 成功返回 1 | 组合: 请求 1, 成功返回 1",
        }, ensure_ascii=False).encode("utf-8"),
    })
    r = parse_zip(z)
    assert r["errors"] == []
    assert len(r["datasets"]) == 1
    assert len(r["metrics"]) == 1
    assert len(r["combinations"]) == 1

    d = r["datasets"][0]
    assert d["name"] == "table1"
    assert d["rowCount"] == 2
    assert d["columns"] == ["MATNR", "LBKUM"]
    # columns.json 的中文注释要挂到字段上
    assert d["fields"][0]["comment"] == "物料编号"
    assert d["fields"][0]["fieldType"] == "VARCHAR(50)"
    assert d["fields"][1]["comment"] == "批次数量"


def test_entries_classified_by_kind():
    z = make_zip({
        "manifest.json": b"{}",
        "tables/t.csv": b"a\n1\n",
        "tables/t.columns.json": b"{}",
        "metrics/m.json": b"{}",
        "combinations/c.json": b"{}",
    })
    from app.data_parser import parse_zip as pz
    r = pz(z)
    kinds = {e["name"]: e["kind"] for e in r["entries"]}
    assert kinds["manifest.json"] == "manifest"
    assert kinds["tables/t.csv"] == "table"
    assert kinds["tables/t.columns.json"] == "columns"
    assert kinds["metrics/m.json"] == "metric"
    assert kinds["combinations/c.json"] == "combination"


def test_classify_helper():
    assert _classify("tables/a.csv") == "table"
    assert _classify("metrics/a.json") == "metric"
    assert _classify("manifest.json") == "manifest"


# ---------- 兼容性 ----------

def test_parse_gbk_csv():
    z = make_zip({"tables/中文表.csv": "姓名,部门\n张三,实验室运营管理部\n".encode("gb18030")})
    r = parse_zip(z)
    assert r["errors"] == []
    assert r["datasets"][0]["preview"][0]["姓名"] == "张三"


def test_parse_flat_csv_without_prefix():
    z = make_zip({"legacy.csv": "a,b\n1,2\n".encode("utf-8-sig")})
    r = parse_zip(z)
    assert len(r["datasets"]) == 1
    assert r["datasets"][0]["name"] == "legacy"


def test_rejects_non_zip():
    with pytest.raises(PlatformParseError):
        parse_zip(b"definitely not a zip")


def test_unknown_type_skipped_not_raised():
    """未知类型静默跳过，但仍列在 entries 里供用户查看。"""
    r = parse_zip(make_zip({"readme.md": b"# hi"}))
    assert r["datasets"] == []
    assert r["errors"] == []
    assert any(e["name"] == "readme.md" for e in r["entries"])


# ---------- 指标结果展开 ----------

def test_flatten_metric_list_payload():
    out = flatten_metric_payload([{"a": 1, "b": "x"}, {"a": 2}])
    assert out["columns"] == ["a", "b"]
    assert len(out["rows"]) == 2


def test_flatten_metric_nested_rows_key():
    out = flatten_metric_payload({"rows": [{"c": 1}]})
    assert out["columns"] == ["c"]


def test_flatten_metric_single_dict():
    out = flatten_metric_payload({"c": 1, "d": 2})
    assert out["columns"] == ["c", "d"]
    assert len(out["rows"]) == 1


# ---------- 字段画像 ----------

def test_profile_numeric():
    p = profile_column(pd.Series(["1", "2", "3.5"], dtype=str))
    assert p["kind"] == "数值"
    assert p["nonEmptyRate"] == 1.0


def test_profile_enum_top():
    p = profile_column(pd.Series(["是"] * 8 + ["否"] * 2, dtype=str))
    assert p["kind"] == "枚举"
    assert p["top"][0] == {"value": "是", "count": 8}


def test_profile_blank_rate():
    p = profile_column(pd.Series(["a", "", "  ", "b"], dtype=str))
    assert p["nonEmpty"] == 2
    assert p["nonEmptyRate"] == 0.5


# ---------- 参数校验（文档 5.3 / 6.3 / 10.10） ----------

def test_page_params_must_come_in_pair():
    with pytest.raises(Exception):
        PullRequest(username="u", pageNo=1)
    with pytest.raises(Exception):
        PullRequest(username="u", pageSize=100)


def test_page_size_bounds():
    with pytest.raises(Exception):
        PullRequest(username="u", pageNo=1, pageSize=10001)
    with pytest.raises(Exception):
        PullRequest(username="u", pageNo=1, pageSize=0)
    with pytest.raises(Exception):
        PullRequest(username="u", pageNo=0, pageSize=100)


def test_page_size_max_allowed():
    assert PullRequest(username="u", pageNo=1, pageSize=10000).pageSize == 10000


def test_condition_field_name_must_be_physical():
    # 文档 10.10：只允许字母数字下划线，且以字母或下划线开头
    for bad in ["1ABC", "有中文", "has-dash", "has space", ""]:
        with pytest.raises(Exception):
            ConditionItem(field=bad, operator="eq", value="1")


def test_condition_field_ok():
    assert ConditionItem(field="MATNR_01", operator="eq", value="1").field == "MATNR_01"
    assert ConditionItem(field="_x", operator="eq", value="1").field == "_x"


def test_condition_operator_enum():
    for op in ["eq", "like", "gte", "lte", "in", "isNull", "isNotNull"]:
        item = ConditionItem(field="F", operator=op, value=None if "Null" in op else "v")
        assert item.operator == op
    with pytest.raises(Exception):
        ConditionItem(field="F", operator="contains", value="v")


def test_condition_in_requires_value():
    with pytest.raises(Exception):
        ConditionItem(field="F", operator="in", value=[])
    assert ConditionItem(field="F", operator="in", value=["1", "2"]).value == ["1", "2"]


def test_condition_isnull_needs_no_value():
    assert ConditionItem(field="F", operator="isNull", value="").value == ""


def test_condition_eq_requires_value():
    with pytest.raises(Exception):
        ConditionItem(field="F", operator="eq", value="")


def test_pass_rate_lifnr_required_and_werk_len():
    with pytest.raises(Exception):
        PassRateRequest(username="u")
    with pytest.raises(Exception):
        PassRateRequest(username="u", lifnr="0001", werk="12345")


# ---------- Token 缓存 ----------

def test_token_cache_decodes_exp():
    # exp = 4102444800（2100-01-01），减 300s 提前量
    token = (
        "eyJhbGciOiJIUzI1NiJ9."
        "eyJleHAiOjQxMDI0NDQ4MDB9."  # {"exp":4102444800}
        "sig"
    )
    exp = TokenCache._decode_exp(token)
    assert exp is not None
    assert 4102444800 - 301 <= exp <= 4102444800 - 299


def test_token_cache_bad_token_returns_none():
    assert TokenCache._decode_exp("not-a-jwt") is None


def test_token_cache_clear_resets():
    c = TokenCache()
    c._token = "x"
    c._expire_at = time.time() + 999
    c.clear()
    assert c._token is None
    assert c._expire_at == 0


# ---------- 固定配置（文档 AI 集成说明） ----------

def test_documented_login_account_is_used():
    from app import config as cfgmod
    cfg = cfgmod.load_config()
    assert cfg["loginUsername"] == "api_data"
    assert cfg["loginPassword"] == "M_bWFjY3VyYTEyMw=="
    assert cfg["apiBase"] == "http://10.19.51.251:8080"


def test_save_config_cannot_override_fixed_values():
    from app import config as cfgmod
    saved = cfgmod.save_config({"username": "someone", "apiBase": "http://evil"})
    assert saved["username"] == "someone"
    # 固定配置必须被强制回落，防止页面改坏
    assert saved["apiBase"] == "http://10.19.51.251:8080"
    assert saved["loginUsername"] == "api_data"
    # 还原
    cfgmod.save_config({"username": "10399"})


def test_zip_ttl_is_three_minutes():
    assert ZIP_TTL_SECONDS == 180


# ---------- 接口层（TestClient，防回归） ----------

@pytest.fixture()
def client(monkeypatch):
    """真实路由 + 打桩的外部调用，不联网。"""
    from fastapi.testclient import TestClient
    import app.main as main_mod

    monkeypatch.setattr(
        main_mod, "TOKEN_CACHE",
        type("TC", (), {"get": staticmethod(lambda cfg: ("stub-token", True)),
                       "clear": staticmethod(lambda: None)})(),
    )
    return TestClient(main_mod.app)


def test_api_meta_returns_operator_enum(client):
    import app.main as main_mod

    r = client.get("/api/meta")
    assert r.status_code == 200
    body = r.json()
    assert body["ok"] is True
    ops = body["data"]["operators"]
    assert {o["value"] for o in ops} == {
        "eq", "like", "gte", "lte", "in", "isNull", "isNotNull"
    }
    # isNull/isNotNull 标记为无需填值，页面据此禁用输入框
    noval = {o["value"] for o in ops if o["noValue"]}
    assert noval == {"isNull", "isNotNull"}
    fields = body["data"]["tableFields"]["ADS_SERVICE_DEMAND_FEEDBACK"]
    # 字段来源：真实拉取缓存优先，缺失时回退内置示例；两种来源都必须有 name/comment/values
    assert fields, "默认表必须有字段元数据"
    for f in fields:
        assert "name" in f and "comment" in f
        assert isinstance(f.get("values", []), list)
    names = {f["name"] for f in fields}
    assert {"YJFL", "ZDKH", "DQ"} <= names or names <= {
        f["name"] for f in main_mod.TABLE_FIELDS["ADS_SERVICE_DEMAND_FEEDBACK"]
    }


def test_field_cache_roundtrip(tmp_path, monkeypatch):
    """拉取后写入的字段缓存应能被 /api/meta 读回，字段与真实表一致。"""
    import app.main as main_mod

    cache_file = tmp_path / "_field_cache.json"
    monkeypatch.setattr(main_mod, "FIELD_CACHE_FILE", cache_file)

    datasets = [{
        "name": "ADS_SERVICE_DEMAND_FEEDBACK",
        "chineseName": "服务需求反馈",
        "fields": [
            {"name": "ZDKH", "comment": "终端客户", "fieldType": "VARCHAR(200)",
             "nonEmpty": 2, "dist": [{"value": "柳林县人民医院", "count": 2}]},
            {"name": "DQ", "comment": "大区", "fieldType": "VARCHAR(200)",
             "nonEmpty": 2, "dist": [{"value": "华东一区", "count": 1},
                                     {"value": "西南区", "count": 1}]},
        ],
    }]
    main_mod._update_field_cache(datasets)

    eff = main_mod._effective_table_fields()
    got = eff["ADS_SERVICE_DEMAND_FEEDBACK"]
    assert [f["name"] for f in got] == ["ZDKH", "DQ"]
    assert got[0]["comment"] == "终端客户"
    assert got[0]["values"] == ["柳林县人民医院"]
    assert got[1]["values"] == ["华东一区", "西南区"]


def test_field_cache_ignores_junk_table_names(tmp_path, monkeypatch):
    """明显是测试用的短表名不应进入字段缓存。"""
    import app.main as main_mod

    monkeypatch.setattr(main_mod, "FIELD_CACHE_FILE", tmp_path / "_fc.json")
    main_mod._update_field_cache([
        {"name": "A", "fields": [{"name": "F", "comment": "", "dist": []}]},
        {"name": "ADS_SERVICE_DEMAND_FEEDBACK", "fields": [
            {"name": "DQ", "comment": "大区", "dist": [{"value": "西南区", "count": 1}]}]},
    ])
    eff = main_mod._effective_table_fields()
    assert "A" not in eff
    assert "ADS_SERVICE_DEMAND_FEEDBACK" in eff


def test_parser_outputs_business_stats():
    """字段画像需含枚举分布、数值统计与日期范围，供业务分布图使用。"""
    from app.data_parser import analyze_dataframe
    import pandas as pd

    df = pd.DataFrame({
        "DQ": ["西南区", "西南区", "华东一区"],
        "SCORE": ["4.5", "3.5", "5"],
        "CREATEDATE": ["2026-10-01", "2026-10-03", "2026-10-02"],
    })
    ds = analyze_dataframe(df, "T1", {"SCORE": {"comment": "得分", "fieldType": "DECIMAL"}})

    by_name = {f["name"]: f for f in ds["fields"]}
    assert by_name["DQ"]["dist"] == [
        {"value": "西南区", "count": 2}, {"value": "华东一区", "count": 1}
    ]
    assert by_name["SCORE"]["numStats"]["avg"] == 4.3333
    assert by_name["SCORE"]["numStats"]["max"] == 5.0
    assert by_name["CREATEDATE"]["dateRange"] == {"min": "2026-10-01", "max": "2026-10-03"}
    assert by_name["SCORE"]["comment"] == "得分"


def test_httpx_client_does_not_inherit_system_proxy():
    """中台在内网，客户端必须 trust_env=False，否则请求被系统代理拦截。"""
    import app.platform_client as pc

    client = pc.DataPlatformClient({"apiBase": "http://10.19.51.251:8080"})
    try:
        assert client._client.trust_env is False
    finally:
        client.close()


def test_api_pull_rejects_condition_key_not_in_tablename(client, monkeypatch):
    """文档 10.7：condition 的 key 必须在 tableName 中，否则 400 而不是 500。"""
    import app.main as main_mod
    called = {"n": 0}

    class Boom(main_mod.DataPlatformClient):
        def __init__(self, cfg):
            called["n"] += 1

        def close(self):
            pass

    monkeypatch.setattr(main_mod, "DataPlatformClient", Boom)
    r = client.post("/api/pull", json={
        "username": "u", "tableName": ["A"],
        "condition": {"B": [{"field": "X", "operator": "eq", "value": "1"}]},
    })
    assert r.status_code == 400
    assert r.json()["ok"] is False
    assert "condition" in r.json()["message"]
    assert called["n"] == 1  # 校验在构造客户端后立即执行，不发请求


def test_api_pull_valid_condition_passes_validation(client, monkeypatch):
    """合法 condition 不该在校验层被拦（后续失败由外部接口决定）。"""
    import app.main as main_mod
    seen = {}

    class FakeClient(main_mod.DataPlatformClient):
        def __init__(self, cfg):
            pass

        def pull_file(self, body, token=None):
            seen["body"] = body
            return {"downloadUrl": "http://x/y.zip", "fileName": "y.zip"}

        def download_zip(self, url):
            return make_zip({
                "tables/A.csv": "F\n1\n".encode("utf-8-sig"),
                "manifest.json": b'{"tables":[],"metrics":[],"combinations":[],"message":"ok"}',
            })

        def close(self):
            pass

    monkeypatch.setattr(main_mod, "DataPlatformClient", FakeClient)
    r = client.post("/api/pull", json={
        "username": "u", "tableName": ["A"],
        "condition": {"A": [{"field": "F", "operator": "eq", "value": "1"}]},
    })
    assert r.status_code == 200
    assert r.json()["ok"] is True
    # condition 正确透传给接口
    assert seen["body"]["condition"] == {"A": [{"field": "F", "operator": "eq", "value": "1"}]}


def test_api_pull_rejects_bad_field_name(client):
    r = client.post("/api/pull", json={
        "username": "u",
        "condition": {"A": [{"field": "1BAD", "operator": "eq", "value": "1"}]},
    })
    assert r.status_code == 422


def test_api_pull_rejects_lone_pageno(client):
    r = client.post("/api/pull", json={"username": "u", "pageNo": 1})
    assert r.status_code == 422


def test_api_sample_returns_full_structure(client):
    from tools.make_sample_zip import build

    build()
    r = client.post("/api/pull/sample")
    assert r.status_code == 200
    d = r.json()["data"]
    assert len(d["datasets"]) == 2
    assert len(d["metrics"]) == 1
    assert len(d["combinations"]) == 1
    # 指标要能展平成表格
    assert d["metrics"][0]["table"]["rows"]
    assert d["metrics"][0]["table"]["columns"]
    # 表字段要带中文注释
    assert any(f["comment"] for t in d["datasets"] for f in t["fields"])


# ---------- 样本 ----------

def test_sample_matches_document_structure():
    from tools.make_sample_zip import build

    p = build()
    r = parse_zip(p.read_bytes())
    assert r["errors"] == []
    assert len(r["datasets"]) == 2
    assert len(r["metrics"]) == 1
    assert len(r["combinations"]) == 1
    # 字段注释要能读出来
    assert any(f["comment"] for d in r["datasets"] for f in d["fields"])
    names = {e["name"] for e in r["entries"]}
    assert "tables/ADS_SERVICE_DEMAND_FEEDBACK.columns.json" in names
    assert "manifest.json" in names


def test_sample_with_paging_has_totals():
    from tools.make_sample_zip import build

    p = build(with_paging=True)
    r = parse_zip(p.read_bytes())
    t = r["manifest"]["tables"][0]
    assert "total" in t and "totalPages" in t
    assert t["pageSize"] == 5

# ---------- BI 看板 ----------

def test_bi_dataset_shape(client):
    """/api/bi/dataset 返回模板可直接消费的中文表头结构。"""
    r = client.get("/api/bi/dataset")
    if r.status_code == 404:
        pytest.skip("尚无真实拉取数据")
    assert r.status_code == 200
    body = r.json()["data"]
    assert body["table"] == "ADS_SERVICE_DEMAND_FEEDBACK"
    for key in ("columns", "rows", "dimensions", "diffAreas"):
        assert key in body
    # 明细行的键必须与 columns 对得上（模板靠中文表头做模糊匹配）
    col_names = {c["key"] for c in body["columns"]}
    if body["rows"]:
        assert set(body["rows"][0].keys()) == col_names


def test_bi_export_xlsx_is_valid_package(client):
    """/api/bi/export.xlsx 必须返回合法 xlsx（ZIP 魔数 PK）。"""
    r = client.get("/api/bi/export.xlsx")
    if r.status_code == 404:
        pytest.skip("尚无真实拉取数据")
    assert r.status_code == 200
    assert r.content[:2] == b"PK"


def test_bi_template_is_offline_capable(client):
    """模板不得再引用公网 CDN，下载版需内联 ECharts。"""
    import app.main as main_mod

    r = client.get("/api/bi/template")
    assert r.status_code == 200
    assert "cdn.jsdelivr.net" not in r.text
    assert "/vendor/echarts.min.js" in r.text

    # 下载版：真实数据已注入 + ECharts 已内联（无外链）
    h = client.get("/api/bi/export.html")
    if h.status_code == 404:
        pytest.skip("尚无真实拉取数据")
    assert "cdn.jsdelivr.net" not in h.text
    assert 'src="/vendor/echarts.min.js"' not in h.text
    assert "ECharts 5.5.0" in h.text


def test_bi_archived_derived_from_workflow_not_empty_column(client):
    """「是否归档」由「问题状态」判断：状态 ∈ 归档状态值即视为已闭环。"""
    from app import bi_mapping

    rows, rules, _ = bi_mapping.load_mapping()
    archived = next(r for r in rows if r["dimension"] == "archived")
    assert archived["method"] == bi_mapping.METHOD_ARCHIVED_BY_STATUS
    # 归档的判据来自问题状态维度，且状态值清单非空
    assert rules.get("归档.状态维度") == "status"
    targets = {v.strip() for v in rules.get("归档.状态值", "").split(",") if v.strip()}
    assert {"临时解决", "完全解决", "无法解决"} <= targets
    # 归档的日期判据不能是 XGPG（真实数据中该列整列为空）
    assert rules.get("归档.日期字段") != "XGPG"

    r = client.get("/api/bi/dataset")
    if r.status_code == 404:
        pytest.skip("尚无真实拉取数据")
    body = r.json()["data"]
    if not body["rows"]:
        pytest.skip("无数据")
    arch = [row["是否归档"] for row in body["rows"]]
    assert set(arch) <= {"是", "否"}
    # 归档口径必须与问题状态一致
    for row in body["rows"]:
        expected = "是" if row.get("问题状态", "") in targets else "否"
        assert row["是否归档"] == expected
    # 口径说明应提到状态判断
    assert any("问题状态" in a and "归档" in a for a in body["diffAreas"])



# ---------- BI 字段映射表（业务方可编辑） ----------

def test_mapping_excel_exists_and_has_sheets():
    """映射表必须存在，且含三个工作表。"""
    from app import bi_mapping

    path = bi_mapping.mapping_path()
    assert path is not None, "缺少 BI字段对应表.xlsx"
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    for name in ("字段对应表", "候选字段清单", "派生规则"):
        assert name in wb.sheetnames, f"缺少工作表 {name}"


def test_mapping_sheet_has_all_dimensions():
    """字段对应表要覆盖全部 18 个模板维度。"""
    from app import bi_mapping

    rows, rules, source = bi_mapping.load_mapping()
    dims = {r["dimension"] for r in rows}
    assert dims == set(bi_mapping.DIMENSION_ORDER)
    # 取数方式必须合法
    for r in rows:
        assert r["method"] in bi_mapping.VALID_METHODS


def test_mapping_change_affects_dataset(tmp_path, monkeypatch):
    """改映射表 → 报表跟着变（业务方改表即生效的核心保证）。"""
    import openpyxl
    from app import bi_mapping

    # 造一份改过的映射表：大区改取 LASTNAME4（解决人）
    src = bi_mapping.mapping_path()
    work = tmp_path / "BI字段对应表.xlsx"
    wb = openpyxl.load_workbook(src)
    ws = wb["字段对应表"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "region":
            row[2].value = "LASTNAME4"
    wb.save(work)

    monkeypatch.setattr(bi_mapping, "_MAPPING_CANDIDATES", [work])
    rows, _, source = bi_mapping.load_mapping()
    assert source == "BI字段对应表.xlsx"
    region = next(r for r in rows if r["dimension"] == "region")
    assert region["field"] == "LASTNAME4"


def test_mapping_endpoint_returns_current_config(client):
    """接口能查到当前生效的映射配置。"""
    r = client.get("/api/bi/mapping")
    assert r.status_code == 200
    data = r.json()["data"]
    assert data["rows"] and data["source"]
    assert "rules" in data


def test_mapping_download_returns_xlsx(client):
    """映射表可下载（业务方要拿到这张表去改）。"""
    r = client.get("/api/bi/mapping.xlsx")
    assert r.status_code == 200
    assert r.content[:2] == b"PK"


# ---------- BI 字段映射表（业务方可编辑） ----------

def test_mapping_excel_exists_and_has_sheets():
    """映射表必须存在，且含三个工作表。"""
    from app import bi_mapping

    path = bi_mapping.mapping_path()
    assert path is not None, "缺少 BI字段对应表.xlsx"
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    for name in ("字段对应表", "候选字段清单", "派生规则"):
        assert name in wb.sheetnames, f"缺少工作表 {name}"


def test_mapping_sheet_has_all_dimensions():
    """字段对应表要覆盖全部 18 个模板维度。"""
    from app import bi_mapping

    rows, rules, source = bi_mapping.load_mapping()
    dims = {r["dimension"] for r in rows}
    assert dims == set(bi_mapping.DIMENSION_ORDER)
    # 取数方式必须合法
    for r in rows:
        assert r["method"] in bi_mapping.VALID_METHODS


def test_mapping_change_affects_dataset(tmp_path, monkeypatch):
    """改映射表 → 报表跟着变（业务方改表即生效的核心保证）。"""
    import openpyxl
    from app import bi_mapping

    # 造一份改过的映射表：大区改取 LASTNAME4（解决人）
    src = bi_mapping.mapping_path()
    work = tmp_path / "BI字段对应表.xlsx"
    wb = openpyxl.load_workbook(src)
    ws = wb["字段对应表"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "region":
            row[2].value = "LASTNAME4"
    wb.save(work)

    monkeypatch.setattr(bi_mapping, "_MAPPING_CANDIDATES", [work])
    rows, _, source = bi_mapping.load_mapping()
    assert source == "BI字段对应表.xlsx"
    region = next(r for r in rows if r["dimension"] == "region")
    assert region["field"] == "LASTNAME4"


def test_mapping_endpoint_returns_current_config(client):
    """接口能查到当前生效的映射配置。"""
    r = client.get("/api/bi/mapping")
    assert r.status_code == 200
    data = r.json()["data"]
    assert data["rows"] and data["source"]
    assert "rules" in data


def test_mapping_download_returns_xlsx(client):
    """映射表可下载（业务方要拿到这张表去改）。"""
    r = client.get("/api/bi/mapping.xlsx")
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
