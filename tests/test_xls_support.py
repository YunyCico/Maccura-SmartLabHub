"""原生工作台（smartlab_engine）对真实 BIFF .xls 的读取回归。

覆盖：工作表列表、原始行读取、表头识别、build_table 建表、导入索引、
大表流式导入，以及「.xls 不允许就地保存覆盖」的写回拦截。
样本由 tools/make_xls_fixtures.py 生成入库，是真正的 OLE/BIFF 文件。
"""
from pathlib import Path

import openpyxl
import pytest

from modules.data_analysis import smartlab_engine as engine

FIXTURES = Path(__file__).resolve().parent / "fixtures"
LIS_SAMPLE = FIXTURES / "lis_sample.xls"
REPORT_SAMPLE = FIXTURES / "report_multilevel_header.xls"


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    lib = tmp_path / "lib"
    (lib / "sources").mkdir(parents=True)
    out = tmp_path / "out"
    out.mkdir()
    monkeypatch.setattr(engine, "LIB_DIR", str(lib))
    monkeypatch.setattr(engine, "OUT_DIR", str(out))
    monkeypatch.setattr(engine, "BAK_DIR", str(out / "_bak"))
    monkeypatch.setattr(engine, "STATE", {"datasets": [], "config": {}})
    engine._cache.clear()
    yield
    engine._cache.clear()


def _sample(path, name):
    target = path / name
    target.write_bytes((LIS_SAMPLE if name == "lis_sample.xls" else REPORT_SAMPLE).read_bytes())
    return target


def test_fixtures_are_real_ole_biff():
    for sample in (LIS_SAMPLE, REPORT_SAMPLE):
        assert sample.is_file(), f"缺少真实 XLS 测试样例：{sample}"
        assert sample.read_bytes()[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", f"{sample.name} 不是 OLE/BIFF 真 .xls"


def test_sheet_names_reads_xls():
    assert engine.sheet_names(str(LIS_SAMPLE)) == ["LIS导出"]
    assert engine.sheet_names(str(REPORT_SAMPLE)) == ["报告"]


def test_cached_raw_reads_xls_rows():
    rows = engine.cached_raw(str(LIS_SAMPLE), "LIS导出")
    assert rows[0][0] == "LIS导出报告"
    assert rows[1][:3] == ["条码号", "姓名", "性别"]
    assert len(rows) == 4
    assert all(len(row) == 7 for row in rows)


def test_build_table_detects_xls_header_and_records():
    rows = engine.cached_raw(str(LIS_SAMPLE), "LIS导出")
    header_rows = engine.detect_header_block(rows, engine.cached_merges(str(LIS_SAMPLE), "LIS导出"))
    table = engine.build_table(str(LIS_SAMPLE), "LIS导出", header_rows)

    assert header_rows == "2"
    assert table["columns"][:3] == ["条码号", "姓名", "性别"]
    assert table["rows"] == 2
    assert table["records"][0][0] == "B001"
    assert table["records"][0][1] == "张三"
    assert table["records"][1][1] == "李四"


def test_xls_and_xlsx_parse_identically(tmp_path):
    """核心验收：同一份数据存成 .xls 与 .xlsx，解析结果必须完全一致。

    这里不写死"应该拼成几级表头"——那是识别启发式的既有行为，
    本次改动的责任是「.xls 与 .xlsx 走同一条识别逻辑，结果不能有差异」。
    """
    import xlwt

    rows = [
        ["基础信息", "基础信息", "检测结果", "检测结果"],
        ["医院名称", "检测完成时间", "模块", "项目名称"],
        ["绵阳市人民医院", "2026-09-28 09:15", "M1", "项目A"],
        ["江油市人民医院", "2026-09-28 10:15", "M2", "项目B"],
    ]

    xlsx_path = tmp_path / "parity.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "报告"
    for row in rows:
        wb.active.append(row)
    wb.save(str(xlsx_path))

    xls_path = tmp_path / "parity.xls"
    book = xlwt.Workbook(encoding="utf-8")
    sheet = book.add_sheet("报告")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            sheet.write(r, c, value)
    book.save(str(xls_path))

    def parse(path):
        names = engine.sheet_names(str(path))
        raw = engine.cached_raw(str(path), names[0])
        header_rows = engine.detect_header_block(raw, engine.cached_merges(str(path), names[0]))
        table = engine.build_table(str(path), names[0], header_rows)
        return names, raw, header_rows, table["columns"], table["rows"], table["records"]

    xls_result = parse(xls_path)
    assert xls_result == parse(xlsx_path), ".xls 与 .xlsx 的解析结果必须完全一致"
    header_rows = xls_result[2]
    assert header_rows in ("1", "1-2"), "表头行识别结果应落在前两行"
    assert xls_result[4] == len(rows) - len(engine.parse_header_spec(header_rows)), "数据行数应与表头行数自洽"


def test_missing_sheet_in_xls_raises():
    with pytest.raises(ValueError, match="工作表不存在"):
        engine.cached_raw(str(LIS_SAMPLE), "不存在")


def test_import_file_indexes_xls(isolated, tmp_path):
    src = _sample(tmp_path, "lis_sample.xls")
    ds = engine.import_file(str(src), copy=True, lab="测试实验室", table_type="LIS")

    assert ds["name"] == "lis_sample.xls"
    assert ds["lab"] == "测试实验室"
    sheet = ds["sheets"][0]
    assert sheet["name"] == "LIS导出"
    assert sheet["header_rows"] == "2"
    assert sheet["rows"] == 2
    assert sheet["columns"][:3] == ["条码号", "姓名", "性别"]
    assert not sheet.get("error"), sheet.get("error")


def test_import_file_rejects_unknown_extension(isolated, tmp_path):
    bad = tmp_path / "data.pdf"
    bad.write_bytes(b"%PDF-1.4")
    with pytest.raises(ValueError, match="暂不支持的格式"):
        engine.import_file(str(bad), copy=False)


def test_large_file_stream_import_reads_xls(isolated, tmp_path, monkeypatch):
    monkeypatch.setattr(engine, "LARGE_FILE_BYTES", 1)
    src = _sample(tmp_path, "lis_sample.xls")
    ds = engine.import_file(str(src), copy=True)

    assert ds["large"] is True
    sheet = ds["sheets"][0]
    assert sheet["large"] is True
    assert sheet["columns"][:3] == ["条码号", "姓名", "性别"]
    page = engine._large_page(ds, sheet, 0, 5)
    assert page[0][0] == "B001"
    assert page[1][0] == "B002"
    assert len(page) == sheet["rows"] == 2


def test_raw_edit_save_blocks_xls_overwrite(isolated, tmp_path):
    src = _sample(tmp_path, "lis_sample.xls")
    ds = engine.import_file(str(src), copy=True)
    before = engine.ds_abs_path(ds)

    with pytest.raises(ValueError, match=r"\.xls 文件暂不支持就地保存"):
        engine.raw_edit_save(ds["id"], sheet="LIS导出", backup=False)

    assert engine.ds_abs_path(ds) == before
    assert open(before, "rb").read()[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "原始 .xls 不得被改写"


def test_raw_edit_save_still_allows_xlsx(isolated, tmp_path):
    path = tmp_path / "edit_ok.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.active.append(["实验室", "值"])
    wb.active.append(["L1", 1])
    wb.save(str(path))

    ds = engine.import_file(str(path), copy=True)
    result = engine.raw_edit_save(ds["id"], sheet="Sheet1", backup=False)

    assert result["ok"] is True
    assert result["path"].endswith(".xlsx")
    assert openpyxl.load_workbook(result["path"]).active.max_row == 2
