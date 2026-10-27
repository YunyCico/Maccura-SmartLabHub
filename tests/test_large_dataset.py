"""大表模式回归：流式导入 SQLite、分页预览、聚合、编辑封锁。全部 tmp 隔离。"""
import openpyxl
import pytest

from modules.data_analysis import smartlab_engine as engine


@pytest.fixture
def large_env(tmp_path, monkeypatch):
    lib = tmp_path / "lib"
    (lib / "sources").mkdir(parents=True)
    out = tmp_path / "out"
    out.mkdir()
    monkeypatch.setattr(engine, "LIB_DIR", str(lib))
    monkeypatch.setattr(engine, "OUT_DIR", str(out))
    monkeypatch.setattr(engine, "BAK_DIR", str(out / "_bak"))
    monkeypatch.setattr(engine, "STATE", {"datasets": [], "config": {}})
    monkeypatch.setattr(engine, "LARGE_FILE_BYTES", 1)   # 任意文件都走大表模式
    engine._cache.clear()
    return lib


def _make(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(rows[0])
    for r in rows[1:]:
        ws.append(r)
    wb.save(str(path))


def test_large_import_sqlite_meta_and_preview(large_env, tmp_path):
    src = tmp_path / "big.xlsx"
    _make(src, [
        ["实验室", "模块", "模块", "值"],
        ["L1", "M1", None, "1"],
        ["L2", "M2", None, "2"],
        ["L3", "M3", None, "3"],
    ])
    ds = engine.import_file(str(src), copy=True)
    assert ds["large"] is True
    sh = ds["sheets"][0]
    assert sh["large"] is True
    import os
    assert os.path.exists(engine._ds_db_path(ds))
    # 重复空表头列被丢，列 = 实验室/模块/值
    assert sh["base_columns"] == ["实验室", "模块", "值"]

    # 分页预览
    page0 = engine._large_page(ds, sh, 0, 2)
    page1 = engine._large_page(ds, sh, 2, 2)
    assert len(page0) == 2 and len(page1) == 1
    assert page0[0][0] == "L1" and page1[0][0] == "L3"


def test_large_aggregate_union_and_cap(large_env, tmp_path, monkeypatch):
    monkeypatch.setattr(engine, "MAX_RESULT_ROWS", 2)   # 验证兜底封顶
    for i, name in enumerate(["a.xlsx", "b.xlsx"]):
        src = tmp_path / name
        _make(src, [
            ["实验室", "模块", "模块", "值"],
            [f"L{i}1", "M1", None, str(i)],
            [f"L{i}2", "M2", None, str(i + 1)],
        ])
        engine.import_file(str(src), copy=True)

    req = {"sources": [{"dataset_id": d["id"], "sheet": "Sheet1"} for d in engine.STATE["datasets"]],
           "fields": ["实验室", "模块", "值"], "add_src": True}
    by_id = {d["id"]: d for d in engine.STATE["datasets"]}
    r = engine.aggregate_union(req, by_id, ["实验室", "模块", "value_x"], [], True, False, {}, 0)
    # 封顶 2 行且带截断标记
    assert len(r["rows"]) == 2 and r["truncated"] is True
    assert r["columns"][0:3] == ["来源实验室", "来源文件", "来源工作表"]
    assert all(x["large"] for x in r["detail"])


def test_large_edit_blocked(large_env, tmp_path):
    src = tmp_path / "big.xlsx"
    _make(src, [["实验室", "模块"], ["L1", "M1"]])
    ds = engine.import_file(str(src), copy=True)
    with pytest.raises(ValueError, match="大表暂不支持"):
        engine.raw_edit_apply(ds["id"], "Sheet1", [{"t": "row", "r": 1}])
    with pytest.raises(ValueError, match="大表暂不支持"):
        engine.raw_edit_save(ds["id"], "Sheet1")


# ---------------------------------------------------------------- 原始数据预览
def test_raw_preview_large_table_endpoint_returns_ok(large_env, tmp_path):
    """大表的「原始数据预览」必须返回 ok 字段。

    回归 bug：/api/raw_preview 的**大表分支**漏了 "ok": True。
    前端 api() 用 `if (!j.ok) throw new Error(j.error || '请求失败（HTTP ' + r.status + '）')`
    判成败 —— 少了 ok，HTTP 明明是 200，界面却报
    「读不到原始数据：请求失败（HTTP 200）」，而且大表必然命中这个分支。
    """
    import json
    import threading
    import urllib.request

    src = tmp_path / "big.xlsx"
    _make(src, [["实验室", "模块", "值"], ["L1", "M1", "1"], ["L2", "M2", "2"]])
    ds = engine.import_file(str(src), copy=True)
    assert ds["large"] is True

    port = engine.pick_port(8799)
    srv = engine.ThreadingHTTPServer(("127.0.0.1", port), engine.Handler)
    srv.daemon_threads = True
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        url = ("http://127.0.0.1:%d/api/raw_preview?dataset_id=%s&sheet=Sheet1"
               "&offset=0&size=10" % (port, ds["id"]))
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        with opener.open(url, timeout=60) as resp:
            body = json.loads(resp.read().decode("utf-8"))
        assert body.get("ok") is True, "大表预览响应必须带 ok 字段，实际：%s" % list(body)[:6]
        assert body.get("large") is True
        assert body.get("total") == 2
        assert len(body.get("rows") or []) == 2
    finally:
        srv.shutdown()


def test_every_engine_json_response_carries_ok():
    """原生引擎的每个 _json 返回体都必须带 ok 字段。

    前端 api() 靠 `ok` 判成败：少了它就会把成功响应当成失败，
    报出"请求失败（HTTP 200）"这种自相矛盾的提示。
    静态扫一遍，新增端点时漏写 ok 会被这条测试挡住。
    """
    from pathlib import Path

    lines = Path(engine.__file__).read_text(encoding="utf-8").split("\n")
    missing = []
    for index, line in enumerate(lines):
        if "self._json({" not in line:
            continue
        window = "\n".join(lines[index:index + 12])
        if '"ok"' not in window:
            missing.append(index + 1)
    assert not missing, "这些 _json 返回体漏了 ok 字段（行号）：%s" % missing


def test_raw_preview_service_supports_large_table(large_env, tmp_path):
    """FastAPI 侧的原始预览也必须支持大表（分页取，不能整表进内存）。"""
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.api.data_analysis import router

    src = tmp_path / "big.xlsx"
    _make(src, [["实验室", "模块", "值"], ["L1", "M1", "1"], ["L2", "M2", "2"]])
    ds = engine.import_file(str(src), copy=True)
    assert ds["large"] is True

    app = FastAPI()
    app.include_router(router, prefix="/api")
    client = TestClient(app)
    response = client.get("/api/data-analysis/raw-preview?dataset_id=%s&sheet=Sheet1&size=10"
                          % ds["id"])
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload.get("large") is True
    assert payload.get("total") == 2
    assert len(payload.get("rows") or []) == 2


# ---------------------------------------------------------------- 原始数据预览
def test_raw_preview_large_table_endpoint_returns_ok(large_env, tmp_path):
    """大表的「原始数据预览」必须返回 ok 字段。

    回归 bug：/api/raw_preview 的**大表分支**漏了 "ok": True。
    前端 api() 用 `if (!j.ok) throw new Error(j.error || '请求失败（HTTP ' + r.status + '）')`
    判成败 —— 少了 ok，HTTP 明明是 200，界面却报
    「读不到原始数据：请求失败（HTTP 200）」，而且大表必然命中这个分支。
    """
    import json
    import threading
    import urllib.request

    src = tmp_path / "big.xlsx"
    _make(src, [["实验室", "模块", "值"], ["L1", "M1", "1"], ["L2", "M2", "2"]])
    ds = engine.import_file(str(src), copy=True)
    assert ds["large"] is True

    port = engine.pick_port(8799)
    srv = engine.ThreadingHTTPServer(("127.0.0.1", port), engine.Handler)
    srv.daemon_threads = True
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        url = ("http://127.0.0.1:%d/api/raw_preview?dataset_id=%s&sheet=Sheet1"
               "&offset=0&size=10" % (port, ds["id"]))
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        with opener.open(url, timeout=60) as resp:
            body = json.loads(resp.read().decode("utf-8"))
        assert body.get("ok") is True, "大表预览响应必须带 ok 字段，实际：%s" % list(body)[:6]
        assert body.get("large") is True
        assert body.get("total") == 2
        assert len(body.get("rows") or []) == 2
    finally:
        srv.shutdown()


def test_every_engine_json_response_carries_ok():
    """原生引擎的每个 _json 返回体都必须带 ok 字段。

    前端 api() 靠 `ok` 判成败：少了它就会把成功响应当成失败，
    报出"请求失败（HTTP 200）"这种自相矛盾的提示。
    静态扫一遍，新增端点时漏写 ok 会被这条测试挡住。
    """
    from pathlib import Path

    lines = Path(engine.__file__).read_text(encoding="utf-8").split("\n")
    missing = []
    for index, line in enumerate(lines):
        if "self._json({" not in line:
            continue
        window = "\n".join(lines[index:index + 12])
        if '"ok"' not in window:
            missing.append(index + 1)
    assert not missing, "这些 _json 返回体漏了 ok 字段（行号）：%s" % missing


def test_raw_preview_service_supports_large_table(large_env, tmp_path):
    """FastAPI 侧的原始预览也必须支持大表（分页取，不能整表进内存）。"""
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.api.data_analysis import router

    src = tmp_path / "big.xlsx"
    _make(src, [["实验室", "模块", "值"], ["L1", "M1", "1"], ["L2", "M2", "2"]])
    ds = engine.import_file(str(src), copy=True)
    assert ds["large"] is True

    app = FastAPI()
    app.include_router(router, prefix="/api")
    client = TestClient(app)
    response = client.get("/api/data-analysis/raw-preview?dataset_id=%s&sheet=Sheet1&size=10"
                          % ds["id"])
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload.get("large") is True
    assert payload.get("total") == 2
    assert len(payload.get("rows") or []) == 2


def test_large_sheet_db_filenames_are_unique(large_env, tmp_path):
    """不同工作表必须写到不同的库文件，且数据不能串。

    回归 bug：旧命名把表名里的中文全换成下划线，于是
    「生化MAX」和「免疫MAX」都变成 "__MAX"、"凝血"和"生化"都变成 "__" ——
    后导入的那个会 os.remove 掉前一个的库文件再重建，
    结果前者查出来是后者的数据（实测索引记 15755 行、库里只有 4250 行）。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "生化MAX"
    ws.append(["项目", "值"])
    ws.append(["生化A", "1"])
    other = wb.create_sheet("免疫MAX")
    other.append(["项目", "值"])
    other.append(["免疫A", "2"])
    other.append(["免疫B", "3"])
    src = tmp_path / "multi.xlsx"
    wb.save(str(src))

    ds = engine.import_file(str(src), copy=True)
    files = [sheet["dbfile"] for sheet in ds["sheets"] if sheet.get("dbfile")]
    assert len(files) == 2, "两个大表工作表应该各有一个库文件"
    assert len(set(files)) == 2, "库文件名撞车了：%s" % files

    # 各自的数据必须是自己的
    for sheet in ds["sheets"]:
        page = engine._large_page(ds, sheet, 0, 10)
        first = page[0][0] if page else ""
        if sheet["name"] == "生化MAX":
            assert first == "生化A", first
        else:
            assert first == "免疫A", first


def test_sheet_db_rel_is_stable_and_unique():
    """库文件命名规则：同表名稳定、不同表名不撞车。"""
    import re as _re

    rel_a = engine._sheet_db_rel("ds_x", "生化MAX")
    rel_b = engine._sheet_db_rel("ds_x", "免疫MAX")
    rel_a2 = engine._sheet_db_rel("ds_x", "生化MAX")
    assert rel_a == rel_a2, "同表名必须生成同一个文件名"
    assert rel_a != rel_b, "不同表名不能撞车"
    # 老规则下这两个会撞成同一个名字
    legacy = lambda n: _re.sub(r"[^A-Za-z0-9_-]", "_", n)
    assert legacy("生化MAX") == legacy("免疫MAX"), "前提：老规则确实会撞车"
