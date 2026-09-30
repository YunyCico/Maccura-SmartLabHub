"""字段删除的回归测试：重复列名（_2 后缀）可删、作用域隔离、源文件真实落盘。

全部在 tmp_path 内构造数据并 monkeypatch 引擎目录，绝不触碰真实数据目录。
"""
import openpyxl

from smartlab_core import smartlab_engine as engine


def _make_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["实验室", "模块", "模块", "值"])   # 重复表头 → 解析为 模块 / 模块_2
    ws.append(["L1", "M1", "M9", 1])
    ws.append(["L2", "M2", "M8", 2])
    wb.save(str(path))


def _isolate(tmp_path, monkeypatch):
    lib = tmp_path / "lib"
    (lib / "sources").mkdir(parents=True)
    out = tmp_path / "out"
    out.mkdir()
    monkeypatch.setattr(engine, "LIB_DIR", str(lib))
    monkeypatch.setattr(engine, "OUT_DIR", str(out))
    monkeypatch.setattr(engine, "BAK_DIR", str(out / "_bak"))
    monkeypatch.setattr(engine, "STATE", {"datasets": [], "config": {}})
    engine._cache.clear()
    return lib


def _base_columns(ds):
    return [c for sh in ds["sheets"] for c in (sh.get("base_columns") or [])]


def test_delete_duplicate_suffixed_field_removes_source_column(tmp_path, monkeypatch):
    _isolate(tmp_path, monkeypatch)
    src = tmp_path / "src.xlsx"
    _make_xlsx(src)
    ds = engine.import_file(str(src), copy=True)

    dup_key = [f["key"] for f in engine.field_dictionary() if f["key"].endswith("2")][0]
    result = engine.delete_fields([dup_key], dataset_ids=[ds["id"]])

    assert result["touched"], "应至少处理一个工作表"
    assert not result["errors"], result["errors"]

    ds = next(d for d in engine.STATE["datasets"] if d["id"] == ds["id"])
    cols = _base_columns(ds)
    assert dup_key not in cols
    assert len(cols) == 3

    # 源文件本身也要少一列（落盘生效）
    wb = openpyxl.load_workbook(engine.ds_abs_path(ds), read_only=True)
    assert wb["Sheet1"].max_column == 3
    wb.close()


def test_delete_is_scoped_and_reports_errors(tmp_path, monkeypatch):
    _isolate(tmp_path, monkeypatch)
    src_a = tmp_path / "a.xlsx"
    src_b = tmp_path / "b.xlsx"
    _make_xlsx(src_a)
    _make_xlsx(src_b)
    ds_a = engine.import_file(str(src_a), copy=True)
    ds_b = engine.import_file(str(src_b), copy=True)

    dup_key = [f["key"] for f in engine.field_dictionary() if f["key"].endswith("2")][0]
    dup_disp = [c for c in _base_columns(ds_a) if str(c).endswith("_2")][0]
    result = engine.delete_fields([dup_key], dataset_ids=[ds_a["id"]])

    assert not result["errors"]
    cols_a = _base_columns(next(d for d in engine.STATE["datasets"] if d["id"] == ds_a["id"]))
    cols_b = _base_columns(next(d for d in engine.STATE["datasets"] if d["id"] == ds_b["id"]))
    assert dup_disp not in cols_a and len(cols_a) == 3
    assert dup_disp in cols_b and len(cols_b) == 4, "未选中的数据集必须保持不变"


def test_delete_unknown_field_is_noop_without_error(tmp_path, monkeypatch):
    _isolate(tmp_path, monkeypatch)
    src = tmp_path / "src.xlsx"
    _make_xlsx(src)
    ds = engine.import_file(str(src), copy=True)
    result = engine.delete_fields(["不存在的字段"], dataset_ids=[ds["id"]])
    assert result["touched"] == []
    assert not result["errors"]
    assert len(_base_columns(next(d for d in engine.STATE["datasets"] if d["id"] == ds["id"]))) == 4
