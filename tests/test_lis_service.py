from __future__ import annotations

from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest

import openpyxl
from fastapi.testclient import TestClient

PROJECT_ROOT = Path(__file__).resolve().parents[1]
BACKEND_DIR = PROJECT_ROOT / "platform" / "backend"
for path in (PROJECT_ROOT, PROJECT_ROOT / "modules", BACKEND_DIR):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from app.services import lis_service
from modules.data_analysis import smartlab_engine


class LisServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        self.old_dirs = (lis_service.STAGING_DIR, lis_service.OUTPUT_DIR, lis_service.CONFIG_DIR)
        lis_service.STAGING_DIR = root / "staging"
        lis_service.OUTPUT_DIR = root / "standardized"
        lis_service.CONFIG_DIR = root / "mappings"
        self.addCleanup(self._restore_dirs)

    def _restore_dirs(self):
        lis_service.STAGING_DIR, lis_service.OUTPUT_DIR, lis_service.CONFIG_DIR = self.old_dirs

    @staticmethod
    def _workbook() -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LIS"
        ws.append(["LIS导出报告"])
        ws.append(["条码号", "姓名", "性别", "检验项目", "检验结果", "接收时间", "报告时间"])
        ws.append(["B001", "张三", "男", "血常规", "正常", "2026-10-01 08:00:00", "2026-10-01 09:00:00"])
        stream = BytesIO()
        wb.save(stream)
        return stream.getvalue()

    def test_preview_imported_dataset_reuses_library_source_and_selected_sheet(self):
        original = self._workbook()
        root = Path(self.temp.name)
        source_dir = root / "library" / "sources"
        source_dir.mkdir(parents=True)
        source_path = source_dir / "ds_test_sample.xlsx"
        source_path.write_bytes(original)
        old_state = smartlab_engine.STATE
        old_lib_dir = smartlab_engine.LIB_DIR
        old_src_dir = smartlab_engine.SRC_DIR
        smartlab_engine.LIB_DIR = str(root / "library")
        smartlab_engine.SRC_DIR = str(source_dir)
        smartlab_engine.STATE = {"datasets": [{
            "id": "ds_test", "name": "已导入 LIS.xlsx", "stored": "sources/ds_test_sample.xlsx",
            "size": len(original), "sheets": [{"name": "LIS", "rows": 1, "columns": ["条码号"]}],
        }]}
        self.addCleanup(lambda: setattr(smartlab_engine, "STATE", old_state))
        self.addCleanup(lambda: setattr(smartlab_engine, "LIB_DIR", old_lib_dir))
        self.addCleanup(lambda: setattr(smartlab_engine, "SRC_DIR", old_src_dir))

        sources = lis_service.imported_sources()
        self.assertEqual([item["id"] for item in sources], ["ds_test"])
        preview = lis_service.preview_imported_dataset("ds_test", "LIS")
        self.assertEqual(preview["filename"], "已导入 LIS.xlsx")
        self.assertEqual(preview["sheet"], "LIS")
        self.assertEqual(preview["row_count"], 1)
        self.assertEqual(source_path.read_bytes(), original)

    def test_imported_dataset_rejects_path_outside_library_sources(self):
        original = self._workbook()
        root = Path(self.temp.name)
        outside = root / "outside.xlsx"
        outside.write_bytes(original)
        old_state = smartlab_engine.STATE
        old_lib_dir = smartlab_engine.LIB_DIR
        old_src_dir = smartlab_engine.SRC_DIR
        smartlab_engine.LIB_DIR = str(root / "library")
        smartlab_engine.SRC_DIR = str(root / "library" / "sources")
        smartlab_engine.STATE = {"datasets": [{"id": "ds_out", "name": "outside.xlsx", "stored": "../outside.xlsx"}]}
        self.addCleanup(lambda: setattr(smartlab_engine, "STATE", old_state))
        self.addCleanup(lambda: setattr(smartlab_engine, "LIB_DIR", old_lib_dir))
        self.addCleanup(lambda: setattr(smartlab_engine, "SRC_DIR", old_src_dir))
        with self.assertRaisesRegex(ValueError, "源文件不存在"):
            lis_service.preview_imported_dataset("ds_out")

    def test_preview_suggests_mapping_and_confirm_saves_standard_copy(self):
        original = self._workbook()
        original_copy = original[:]
        preview = lis_service.preview_upload("sample.xlsx", original)
        self.assertEqual(preview["header_row"], 2)
        # v1.2 起新增 test_specialty / test_suite 两个标准字段，标准字段集共 46 项
        self.assertEqual(len(preview["fields"]), 46)
        self.assertEqual(preview["row_count"], 1)
        self.assertEqual(len(preview["preview"]), 1)

        saved = lis_service.confirm_save(preview["session_id"], preview["mapping"], preview["header_row"], "实验室A")
        output_path, output_name = lis_service.output_file(saved["batch_id"])
        self.assertTrue(output_path.is_file())
        self.assertEqual(output_name, saved["file"])
        output = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
        try:
            rows = list(output["lis_std"].iter_rows(values_only=True))
        finally:
            output.close()
        self.assertEqual(rows[0][0], "记录唯一号")
        # v1.2 起导出只保留有值的列，下标会变，按列名取值
        header = list(rows[0])
        self.assertIn("标本条码号", header)
        self.assertIn("患者姓名", header)
        self.assertEqual(rows[1][header.index("标本条码号")], "B001")
        self.assertEqual(rows[1][header.index("患者姓名")], "张三")
        self.assertEqual(rows[1][header.index("患者性别")], "男")
        self.assertEqual(len(rows), 2)
        self.assertEqual((lis_service.STAGING_DIR / f"{preview['session_id']}.xlsx").exists(), False)
        self.assertEqual(original, original_copy)

    def test_real_xls_preview_and_confirm_save_keep_source_unchanged(self):
        source = Path(__file__).resolve().parent / "fixtures" / "lis_sample.xls"
        self.assertTrue(source.is_file(), f"缺少真实 XLS 测试样例：{source}")
        content = source.read_bytes()
        self.assertEqual(content[:8], b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "样例必须是 OLE/BIFF 真 .xls")

        preview = lis_service.preview_upload(source.name, content)

        self.assertEqual(preview["header_row"], 2)
        self.assertEqual(preview["sheet"], "LIS导出")
        self.assertEqual(preview["row_count"], 2)
        self.assertEqual(
            preview["headers"],
            ["条码号", "姓名", "性别", "检验项目", "检验结果", "接收时间", "报告时间"],
        )

        original = content[:]
        saved = lis_service.confirm_save(preview["session_id"], preview["mapping"], preview["header_row"], "实验室A")

        output_path, output_name = lis_service.output_file(saved["batch_id"])
        self.assertTrue(output_path.is_file())
        self.assertTrue(output_name.endswith(".xlsx"), "标准化结果仍输出 xlsx")
        self.assertEqual(saved["rows"], 2)
        output = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
        try:
            rows = list(output["lis_std"].iter_rows(values_only=True))
        finally:
            output.close()
        self.assertEqual(rows[0][0], "记录唯一号")
        header = list(rows[0])
        self.assertEqual(rows[1][header.index("标本条码号")], "B001")
        self.assertEqual(rows[2][header.index("患者姓名")], "李四")
        self.assertEqual((lis_service.STAGING_DIR / f"{preview['session_id']}.xls").exists(), False)
        self.assertEqual(content, original, "原始 .xls 不得被修改")

    def test_xls_header_row_can_be_reselected(self):
        source = Path(__file__).resolve().parent / "fixtures" / "lis_sample.xls"
        self.assertTrue(source.is_file())

        preview = lis_service.preview_upload(source.name, source.read_bytes(), header_row=1)

        self.assertEqual(preview["header_row"], 1)
        self.assertEqual(preview["row_count"], 3)
        self.assertEqual(preview["headers"][0], "LIS导出报告")
        self.assertEqual(preview["preview"][0][0], "条码号")

    def test_unrecognizable_binary_is_rejected(self):
        """既不是 Excel 也不是文本的文件应被拒绝，且给出可读原因。"""
        with self.assertRaisesRegex(ValueError, "损坏|无法识别|不是有效"):
            lis_service.preview_upload("broken.xls", bytes(range(256)) * 20)

    def test_missing_mapping_no_longer_blocks_save(self):
        """v1.2 起取消必填：源表没有的字段留空，不再拒绝整行保存。"""
        preview = lis_service.preview_upload("sample.xlsx", self._workbook())
        mapping = [item for item in preview["mapping"] if item["target_key"] != "barcode"]
        # 不应再抛 ValueError，缺失字段按留空处理
        saved = lis_service.confirm_save(preview["session_id"], mapping, preview["header_row"])
        output_path, _ = lis_service.output_file(saved["batch_id"])
        self.assertTrue(output_path.is_file())

    def test_csv_accepts_gb18030_and_normalizes_date_and_negative_numeric(self):
        content = "条码号,接收时间,TAT\nB002,20260701,-5分钟\n".encode("gb18030")
        preview = lis_service.preview_upload("source.csv", content)
        self.assertEqual(preview["header_row"], 1)
        self.assertEqual(preview["row_count"], 1)

    def test_empty_columns_are_hidden_in_output(self):
        """v1.2：整列全空的标准字段不体现，但溯源与分类列始终保留。"""
        preview = lis_service.preview_upload("sample.xlsx", self._workbook())
        saved = lis_service.confirm_save(preview["session_id"], preview["mapping"], preview["header_row"])
        output_path, _ = lis_service.output_file(saved["batch_id"])
        output = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
        try:
            header = list(next(output["lis_std"].iter_rows(values_only=True)))
        finally:
            output.close()
        total = len(lis_service._field_definitions())
        self.assertLess(len(header), total, "全空列应被省略")
        self.assertEqual(len(header), saved["column_count"])
        self.assertEqual(len(header), total - saved["hidden_empty_columns"])
        # 该样例只有 7 个源列有值，专业/组套自动分类后有值，两者都不该被隐藏
        self.assertIn("专业", header)
        self.assertIn("组套", header)
        self.assertIn("来源文件", header)
        # 源表完全没有的字段应当消失
        self.assertNotIn("临床诊断", header)

    def test_project_is_classified_into_specialty_and_suite(self):
        """检测项目名称 / 项目英文名称对照试剂分类表填出专业与组套。"""
        self.assertEqual(lis_service.classify_project("白蛋白"), {"test_specialty": "生化", "test_suite": "肝功类"})
        self.assertEqual(lis_service.classify_project("", "TSH"), {"test_specialty": "免疫", "test_suite": "甲状腺功能类"})
        # 分类表里写全称、LIS 里写简称，也要能命中
        self.assertEqual(lis_service.classify_project("D-二聚体")["test_suite"], "心肌标志物类")
        # 表里没有的项目留空，不硬凑
        self.assertEqual(lis_service.classify_project("不存在的项目XYZ"), {"test_specialty": "", "test_suite": ""})
        taxonomy = lis_service._load_taxonomy()
        self.assertGreater(len(taxonomy["name"]) + len(taxonomy["en"]), 300, "分类表应已加载")

    def test_english_name_hit_works_when_chinese_name_differs(self):
        """中文名对不上、英文简称对得上时，也应按英文命中分类。

        分类表是试剂维度，品名写「丙氨酸氨基转移酶（丙氨酸底物法）」，
        LIS 检测项目写科室惯用名「谷丙转氨酶」——英文简称 ALT 才是通用桥梁。
        """
        self.assertEqual(
            lis_service.classify_project("谷丙转氨酶", "ALT"),
            {"test_specialty": "生化", "test_suite": "肝功类"},
        )
        # 英文名直接对上也算命中（用户口径：英文名对得上就是同一个项目）
        self.assertEqual(
            lis_service.classify_project("随便写的名字", "D-Dimer"),
            {"test_specialty": "免疫", "test_suite": "心肌标志物类"},
        )
        # 英文缩写写在中文名括号里，也能命中
        self.assertEqual(
            lis_service.classify_project("谷丙转氨酶(ALT)", "")["test_specialty"],
            "生化",
        )
        # 中文惯用名没有英文列时，靠别名桥同样命中
        self.assertEqual(
            lis_service.classify_project("谷丙转氨酶", "")["test_suite"],
            "肝功类",
        )

    def test_project_name_is_never_rewritten(self):
        """分类只产出专业/组套两列，检测项目名称与项目英文名称必须原样保留。"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LIS"
        ws.append(["条码号", "姓名", "性别", "检验项目", "项目英文名称", "接收时间"])
        ws.append(["B010", "王五", "男", "谷丙转氨酶", "ALT", "2026-10-01 08:00:00"])
        stream = BytesIO()
        wb.save(stream)
        preview = lis_service.preview_upload("keepname.xlsx", stream.getvalue())
        saved = lis_service.confirm_save(preview["session_id"], preview["mapping"], preview["header_row"])
        output_path, _ = lis_service.output_file(saved["batch_id"])
        output = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
        try:
            rows = list(output["lis_std"].iter_rows(values_only=True))
        finally:
            output.close()
        header = list(rows[0])
        self.assertEqual(rows[1][header.index("检测项目名称")], "谷丙转氨酶")
        self.assertEqual(rows[1][header.index("项目英文名称")], "ALT")
        # 分类结果另行给出
        self.assertEqual(rows[1][header.index("专业")], "生化")
        self.assertEqual(rows[1][header.index("组套")], "肝功类")

    def test_suite_level_and_composite_names_are_classified(self):
        """组套级送检项目名、复合项目名也要能分类。"""
        # 组套级别：肝功能 → 生化/肝功类
        self.assertEqual(
            lis_service.classify_project("肝功能", ""),
            {"test_specialty": "生化", "test_suite": "肝功类"},
        )
        self.assertEqual(
            lis_service.classify_project("肾功能", ""),
            {"test_specialty": "生化", "test_suite": "肾功类"},
        )
        # 复合项目名：逐段尝试，命中任一段即用
        result = lis_service.classify_project("ASO、RF、CCP", "")
        self.assertTrue(result["test_specialty"], "复合项目名应至少命中一段")
        # 剥修饰词：血清钾测定 → 钾
        self.assertEqual(
            lis_service.classify_project("钙测定", ""),
            {"test_specialty": "生化", "test_suite": "电解质类"},
        )

    def test_classify_api_endpoint(self):
        """POST /data-analysis/lis/classify 批量判定，规则与后端一致。"""
        from app.api.data_analysis import router
        from fastapi import FastAPI

        app = FastAPI()
        app.include_router(router)
        client = TestClient(app)
        response = client.post(
            "/data-analysis/lis/classify",
            json={
                "rows": [
                    {"test_name": "谷丙转氨酶", "test_en_name": "ALT"},
                    {"test_name": "肝功能", "test_en_name": ""},
                    {"test_name": "不存在的项目XYZ", "test_en_name": ""},
                ]
            },
        )
        self.assertEqual(response.status_code, 200)
        results = response.json()["results"]
        self.assertEqual(len(results), 3)
        self.assertEqual(results[0]["test_specialty"], "生化")
        self.assertEqual(results[0]["test_suite"], "肝功类")
        self.assertEqual(results[1]["test_suite"], "肝功类")
        self.assertEqual(results[2]["test_specialty"], "")

    def test_saved_rows_carry_classification(self):
        """导出的行里专业/组套列必须带上分类结果。"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LIS"
        ws.append(["条码号", "姓名", "性别", "检验项目", "项目英文名称", "接收时间", "报告时间"])
        ws.append(["B009", "李四", "女", "白蛋白", "ALB", "2026-10-01 08:00:00", "2026-10-01 09:00:00"])
        stream = BytesIO()
        wb.save(stream)
        preview = lis_service.preview_upload("classify.xlsx", stream.getvalue())
        saved = lis_service.confirm_save(preview["session_id"], preview["mapping"], preview["header_row"])
        output_path, _ = lis_service.output_file(saved["batch_id"])
        output = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
        try:
            rows = list(output["lis_std"].iter_rows(values_only=True))
        finally:
            output.close()
        header = list(rows[0])
        self.assertEqual(rows[1][header.index("专业")], "生化")
        self.assertEqual(rows[1][header.index("组套")], "肝功类")

    def test_user_confirmed_taxonomy_additions(self):
        """2026-10-20 用户在核对表 D 列核实的 7 个品名必须能分类。

        口径：品名在分类表里 → 分类表加行；品名已存在 → 只做名称翻译。
        专业/组套一律回分类表查，代码里不硬编码。
        """
        cases = [
            ("血清钾测定", "生化", "电解质类"),
            ("血清钠测定", "生化", "电解质类"),
            ("血清氯测定", "生化", "电解质类"),
            ("白带检查", "临检", "体液类"),
            ("血气分析", "生化", "血气类"),
            ("白细胞介素6", "免疫", "炎症类"),
            ("粪便隐血", "临检", "体液类"),
        ]
        for test_name, specialty, suite in cases:
            with self.subTest(test_name=test_name):
                result = lis_service.classify_project(test_name, "")
                self.assertEqual(result["test_specialty"], specialty)
                self.assertEqual(result["test_suite"], suite)

    def test_taxonomy_additions_exist_in_classification_table(self):
        """新增品名必须真的落在分类表里（而不是只靠代码里的映射蒙对）。"""
        items = lis_service.taxonomy_items()
        by_name = {item["name"]: item for item in items}
        for name, specialty, suite in (
            ("钾", "生化", "电解质类"),
            ("钠", "生化", "电解质类"),
            ("氯", "生化", "电解质类"),
            ("血气", "生化", "血气类"),
            ("白带", "临检", "体液类"),
            ("便隐血", "临检", "体液类"),
        ):
            with self.subTest(name=name):
                self.assertIn(name, by_name)
                self.assertEqual(by_name[name]["specialty"], specialty)
                self.assertEqual(by_name[name]["suite"], suite)

    def test_added_names_do_not_rewrite_lis_project_name(self):
        """加行只影响专业/组套，检测项目名称与项目英文名称必须原样保留。"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LIS"
        ws.append(["条码号", "姓名", "性别", "检验项目", "项目英文名称", "接收时间", "报告时间"])
        ws.append(["B010", "王五", "女", "血气分析", "ABG", "2026-10-01 08:00:00", "2026-10-01 09:00:00"])
        stream = BytesIO()
        wb.save(stream)
        preview = lis_service.preview_upload("gas.xlsx", stream.getvalue())
        saved = lis_service.confirm_save(preview["session_id"], preview["mapping"], preview["header_row"])
        output_path, _ = lis_service.output_file(saved["batch_id"])
        output = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
        try:
            rows = list(output["lis_std"].iter_rows(values_only=True))
        finally:
            output.close()
        header = list(rows[0])
        self.assertEqual(rows[1][header.index("检测项目名称")], "血气分析")
        self.assertEqual(rows[1][header.index("项目英文名称")], "ABG")
        self.assertEqual(rows[1][header.index("专业")], "生化")
        self.assertEqual(rows[1][header.index("组套")], "血气类")


class LisPreviewFlowTests(unittest.TestCase):
    """「确认字段对应 → 出预览 → 导出 / 保存」这条新链路。"""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        for name in ("staging", "standardized", "mappings"):
            (root / name).mkdir(parents=True, exist_ok=True)
        self.old_dirs = (lis_service.STAGING_DIR, lis_service.OUTPUT_DIR, lis_service.CONFIG_DIR)
        lis_service.STAGING_DIR = root / "staging"
        lis_service.OUTPUT_DIR = root / "standardized"
        lis_service.CONFIG_DIR = root / "mappings"
        lis_service._PAYLOAD_CACHE.clear()
        self.addCleanup(self._restore)

    def _restore(self):
        (lis_service.STAGING_DIR, lis_service.OUTPUT_DIR, lis_service.CONFIG_DIR) = self.old_dirs
        lis_service._PAYLOAD_CACHE.clear()

    def _workbook(self, rows: int = 5) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LIS"
        ws.append(["LIS导出报告"])
        ws.append(["条码号", "姓名", "性别", "检验项目", "项目英文名称", "接收时间"])
        for index in range(rows):
            ws.append([
                f"B{index:03d}", f"患者{index}", "男", "谷丙转氨酶", "ALT",
                "2026-10-01 08:00:00",
            ])
        stream = BytesIO()
        wb.save(stream)
        return stream.getvalue()

    def test_preview_standard_returns_all_rows_across_pages(self):
        """预览要能翻页看完所有数据行，而不是只给前 N 行。"""
        preview = lis_service.preview_upload("sample.xlsx", self._workbook(5))
        first = lis_service.preview_standard(
            preview["session_id"], preview["mapping"], preview["header_row"], "", page=1, page_size=2)
        self.assertEqual(first["total_rows"], 5)
        self.assertEqual(first["total_pages"], 3)
        self.assertEqual(len(first["rows"]), 2)
        self.assertEqual(first["page"], 1)
        second = lis_service.preview_standard(
            preview["session_id"], preview["mapping"], preview["header_row"], "", page=2, page_size=2)
        self.assertEqual(len(second["rows"]), 2)
        third = lis_service.preview_standard(
            preview["session_id"], preview["mapping"], preview["header_row"], "", page=3, page_size=2)
        self.assertEqual(len(third["rows"]), 1)
        # 三页拼起来正好是全部 5 行，不重不漏
        seen = [row[first["columns"].index("标本条码号")] for row in first["rows"] + second["rows"] + third["rows"]]
        self.assertEqual(seen, ["B000", "B001", "B002", "B003", "B004"])

    def test_preview_standard_matches_saved_output(self):
        """预览所见即保存所得 —— 同一份映射，列和行必须完全一致。"""
        preview = lis_service.preview_upload("sample.xlsx", self._workbook(3))
        standard = lis_service.preview_standard(
            preview["session_id"], preview["mapping"], preview["header_row"], "", page=1, page_size=50)
        saved = lis_service.confirm_save(
            preview["session_id"], preview["mapping"], preview["header_row"], "实验室A")
        output_path, _ = lis_service.output_file(saved["batch_id"])
        output = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
        try:
            rows = list(output["lis_std"].iter_rows(values_only=True))
        finally:
            output.close()
        self.assertEqual(list(rows[0]), standard["columns"])
        self.assertEqual([list(row) for row in rows[1:]], standard["rows"])

    def test_export_writes_excel_without_recording_saved_batch(self):
        """导出只出文件，不进"已整理"列表（那是保存才做的事）。"""
        before = {item["batch_id"] for item in lis_service.saved_batches()}
        preview = lis_service.preview_upload("sample.xlsx", self._workbook(2))
        path, name = lis_service.export_standard(
            preview["session_id"], preview["mapping"], preview["header_row"], "")
        self.assertTrue(path.is_file())
        self.assertTrue(name.startswith("LIS标准化导出_"))
        after = {item["batch_id"] for item in lis_service.saved_batches()}
        self.assertEqual(before, after)

    def test_saved_batches_lists_saved_results(self):
        """保存后应出现在"已整理且保存的数据表格"列表里。"""
        preview = lis_service.preview_upload("sample.xlsx", self._workbook(2))
        saved = lis_service.confirm_save(
            preview["session_id"], preview["mapping"], preview["header_row"], "佳佳")
        batches = lis_service.saved_batches()
        self.assertTrue(any(item["batch_id"] == saved["batch_id"] for item in batches))
        item = next(item for item in batches if item["batch_id"] == saved["batch_id"])
        self.assertEqual(item["source_filename"], "sample.xlsx")
        self.assertEqual(item["lab"], "佳佳")
        self.assertEqual(item["row_count"], 2)
        self.assertTrue(item["available"])

    def test_changing_mapping_invalidates_cached_preview(self):
        """改了映射必须重新构建，不能拿旧映射的预览糊弄用户。"""
        preview = lis_service.preview_upload("sample.xlsx", self._workbook(2))
        first = lis_service.preview_standard(
            preview["session_id"], preview["mapping"], preview["header_row"], "", page=1, page_size=50)
        self.assertIn("检测项目名称", first["columns"])
        changed = [dict(item) for item in preview["mapping"]]
        for item in changed:
            if item["source"] == "检验项目":
                item["target_key"] = ""
                item["target_label"] = ""
        second = lis_service.preview_standard(
            preview["session_id"], changed, preview["header_row"], "", page=1, page_size=50)
        self.assertNotIn("检测项目名称", second["columns"])

    def test_preview_standard_api_endpoint(self):
        """POST /data-analysis/lis/preview-standard 走通。"""
        from app.api.data_analysis import router
        from fastapi import FastAPI

        app = FastAPI()
        app.include_router(router)
        client = TestClient(app)
        preview = lis_service.preview_upload("sample.xlsx", self._workbook(4))
        response = client.post(
            "/data-analysis/lis/preview-standard",
            json={
                "session_id": preview["session_id"],
                "mapping": preview["mapping"],
                "header_row": preview["header_row"],
                "lab": "",
                "page": 1,
                "page_size": 3,
            },
        )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["total_rows"], 4)
        self.assertEqual(body["total_pages"], 2)
        self.assertEqual(len(body["rows"]), 3)

    def test_saved_batches_api_endpoint(self):
        """GET /data-analysis/lis/saved-batches 走通。"""
        from app.api.data_analysis import router
        from fastapi import FastAPI

        app = FastAPI()
        app.include_router(router)
        client = TestClient(app)
        preview = lis_service.preview_upload("sample.xlsx", self._workbook(1))
        lis_service.confirm_save(preview["session_id"], preview["mapping"], preview["header_row"], "")
        response = client.get("/data-analysis/lis/saved-batches")
        self.assertEqual(response.status_code, 200)
        self.assertGreaterEqual(len(response.json()["batches"]), 1)

    def test_delete_saved_batch_removes_record_and_file(self):
        """删除一条已整理记录：映射 json 和标准化 Excel 都要消失，列表里也不再出现。"""
        preview = lis_service.preview_upload("sample.xlsx", self._workbook(2))
        saved = lis_service.confirm_save(
            preview["session_id"], preview["mapping"], preview["header_row"], "佳佳")
        batch_id = saved["batch_id"]
        output_path = lis_service.OUTPUT_DIR / saved["file"]
        self.assertTrue(output_path.is_file())
        self.assertTrue((lis_service.CONFIG_DIR / f"{batch_id}.json").is_file())

        result = lis_service.delete_saved_batch(batch_id)
        self.assertTrue(result["deleted_file"])
        self.assertFalse(output_path.is_file())
        self.assertFalse((lis_service.CONFIG_DIR / f"{batch_id}.json").is_file())
        self.assertFalse(any(item["batch_id"] == batch_id for item in lis_service.saved_batches()))

    def test_delete_saved_batch_keeps_other_records(self):
        """删一条不能误伤别的记录。"""
        first = lis_service.preview_upload("a.xlsx", self._workbook(1))
        saved_a = lis_service.confirm_save(
            first["session_id"], first["mapping"], first["header_row"], "")
        second = lis_service.preview_upload("b.xlsx", self._workbook(1))
        saved_b = lis_service.confirm_save(
            second["session_id"], second["mapping"], second["header_row"], "")

        lis_service.delete_saved_batch(saved_a["batch_id"])
        remaining = {item["batch_id"] for item in lis_service.saved_batches()}
        self.assertNotIn(saved_a["batch_id"], remaining)
        self.assertIn(saved_b["batch_id"], remaining)
        self.assertTrue((lis_service.OUTPUT_DIR / saved_b["file"]).is_file())

    def test_delete_saved_batch_rejects_invalid_id(self):
        """批次号必须严格是 12 位十六进制，防止路径穿越。"""
        for bad in ("../../etc/passwd", "abc", "", "ZZZZZZZZZZZZ", "a" * 13):
            with self.assertRaises(Exception):
                lis_service.delete_saved_batch(bad)

    def test_delete_saved_batch_missing_record_raises(self):
        """记录不存在时明确报错，而不是静默成功。"""
        with self.assertRaises(FileNotFoundError):
            lis_service.delete_saved_batch("0123456789ab")

    def test_delete_saved_batch_api_endpoint(self):
        """DELETE /data-analysis/lis/saved-batches/{batch_id} 走通。"""
        from app.api.data_analysis import router
        from fastapi import FastAPI

        app = FastAPI()
        app.include_router(router)
        client = TestClient(app)
        preview = lis_service.preview_upload("sample.xlsx", self._workbook(1))
        saved = lis_service.confirm_save(
            preview["session_id"], preview["mapping"], preview["header_row"], "")
        response = client.delete(f"/data-analysis/lis/saved-batches/{saved['batch_id']}")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["deleted_file"])
        self.assertFalse(any(
            item["batch_id"] == saved["batch_id"]
            for item in client.get("/data-analysis/lis/saved-batches").json()["batches"]))

    def test_delete_saved_batch_api_endpoint_404(self):
        """删不存在的批次返回 404。"""
        from app.api.data_analysis import router
        from fastapi import FastAPI

        app = FastAPI()
        app.include_router(router)
        client = TestClient(app)
        response = client.delete("/data-analysis/lis/saved-batches/0123456789ab")
        self.assertEqual(response.status_code, 404)


    def _multi_sheet_workbook(self) -> bytes:
        """造一个多工作表工作簿，专门用来钉死"读错工作表"这个坑。

        · 「列多的表」：8 列，没有「采样-报告」——自动挑选规则会选它
        · 「用户选的那张」：4 列，有「采样-报告」——用户真正要的
        """
        wb = openpyxl.Workbook()
        wide = wb.active
        wide.title = "列多的表"
        wide.append(["条码号", "姓名", "性别", "年龄", "检验项目", "检验结果", "接收时间", "报告时间"])
        wide.append(["B001", "张三", "男", "41", "血常规", "正常",
                     "2026-10-01 08:00:00", "2026-10-01 09:00:00"])
        narrow = wb.create_sheet("用户选的那张")
        narrow.append(["条码号", "姓名", "检验项目", "采样-报告"])
        narrow.append(["B002", "李四", "血常规", "35"])
        stream = BytesIO()
        wb.save(stream)
        return stream.getvalue()

    def test_standard_payload_keeps_the_sheet_selected_at_preview(self):
        """确认阶段必须沿用「预览时选中的那张工作表」。

        回归 bug：_standard_payload 读暂存文件时没把 sheet 名传进去，
        _read_workbook 就会把工作簿里所有表都读一遍、再按"列数多、行数多"
        自动挑一个 —— 挑到的往往不是用户选的那张，于是字段映射里的来源列
        在表头里找不到，报「来源字段不存在：xxx」。
        实测《标本明细》：用户选「生化」，自动挑中的是「2026.07数据」，
        前者的 TAT 列叫「采样-报告」、后者叫「采集-报告」，就差一个字就报错。
        """
        preview = lis_service.preview_upload(
            "multi.xlsx", self._multi_sheet_workbook(), None, "用户选的那张")
        self.assertEqual(preview["sheet"], "用户选的那张")
        self.assertIn("采样-报告", preview["headers"])
        result = lis_service.preview_standard(
            preview["session_id"], preview["mapping"], preview["header_row"], "")
        self.assertEqual(result["total_rows"], 1)

    def test_standard_payload_rejects_source_column_from_another_sheet(self):
        """反向确认：映射里塞了别的表才有的列，仍然必须报「来源字段不存在」。"""
        preview = lis_service.preview_upload(
            "multi.xlsx", self._multi_sheet_workbook(), None, "用户选的那张")
        mapping = list(preview["mapping"]) + [{
            "source": "性别", "target_key": "", "target_label": "", "rule": "原值保留",
        }]
        # 「性别」只在「列多的表」里，用户选的那张没有 → 必须报错而不是静默忽略
        with self.assertRaises(ValueError):
            lis_service.preview_standard(
                preview["session_id"], mapping, preview["header_row"], "")


if __name__ == "__main__":
    unittest.main()
